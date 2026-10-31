"""Import BRACS files in place, retaining every supplied metadata row."""

from __future__ import annotations

import csv
from datetime import date, datetime
import hashlib
import json
import math
from pathlib import Path
import re

from ..models import AcquisitionPlan, Asset, MetadataRecord, utc_now


SOURCE_URL = "https://www.bracs.icar.cnr.it/"
_ALIASES = {
    "slide_id": ("WSI", "WSI ID", "WSI name", "WSI filename"),
    "label": ("WSI label",),
    "subject_id": ("Patient ID",),
    "split": ("Set", "Split", "Reference set", "WSI set"),
}


def _key(value):
    return re.sub(r"[\s_]+", " ", str(value).strip()).casefold()


def _text(value):
    return "" if value is None else str(value).strip()


def _clean(value):
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, dict):
        return {str(key): _clean(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_clean(item) for item in value]
    if hasattr(value, "item"):
        return _clean(value.item())
    return value


def _read_metadata(path, sheet=None):
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            headers = reader.fieldnames
            rows = list(reader)
        if any(None in row for row in rows):
            raise ValueError("BRACS CSV has rows wider than its header")
    elif path.suffix.lower() == ".xlsx":
        try:
            import openpyxl
        except ImportError as error:
            raise ValueError(
                "Install openpyxl to read BRACS XLSX metadata, or supply a CSV export"
            ) from error
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet is None:
                if "WSI_Information" in book.sheetnames:
                    sheet = "WSI_Information"
                elif len(book.sheetnames) == 1:
                    sheet = book.sheetnames[0]
                else:
                    raise ValueError("Choose the WSI worksheet with metadata['sheet']")
            if sheet not in book.sheetnames:
                raise ValueError("Selected BRACS metadata worksheet does not exist")
            values = iter(book[sheet].iter_rows(values_only=True))
            headers = next(values, None)
            rows = (
                [
                    dict(zip(headers, cells))
                    for cells in values
                    if any(cell is not None for cell in cells)
                ]
                if headers
                else []
            )
        finally:
            book.close()
    else:
        raise ValueError("BRACS metadata must be the supplied XLSX workbook or a CSV export")
    if not headers:
        raise ValueError("BRACS metadata has no header")
    normalized = [_key(header) for header in headers]
    if any(not key or key == "none" for key in normalized) or len(set(normalized)) != len(
        normalized
    ):
        raise ValueError("BRACS metadata needs nonempty, unique column names")
    return _clean(rows), dict(zip(normalized, headers)), sheet


def _resolve_columns(headers, mapping):
    mapping = mapping or {}
    unknown = set(mapping) - set(_ALIASES)
    if unknown:
        raise ValueError("column_mapping keys must be slide_id, label, subject_id, or split")
    columns = {}
    for name, aliases in _ALIASES.items():
        if name in mapping:
            actual = headers.get(_key(mapping[name]))
            if actual is None:
                raise ValueError(f"Mapped BRACS column is absent: {name}")
            columns[name] = actual
        else:
            matches = [headers[_key(alias)] for alias in aliases if _key(alias) in headers]
            if len(set(matches)) > 1:
                raise ValueError(f"Several BRACS columns match {name}; provide column_mapping")
            columns[name] = matches[0] if matches else None
    if columns["slide_id"] is None:
        raise ValueError(
            "BRACS requires an explicit WSI identifier column; provide column_mapping['slide_id']"
        )
    return columns


def _identity(value):
    return hashlib.sha256(value.encode("utf-8")).hexdigest()[:24]


def import_local(root, metadata) -> AcquisitionPlan:
    """Register SVS files inside root; do not copy ROIs or access login tools.

    ``metadata`` is a CSV/XLSX path or a dict containing ``path``, optional
    ``sheet``, ``column_mapping``, ``slides_dir``, and ``glob``. Mappings use the
    canonical keys slide_id, label, subject_id, split. Unknown source columns
    are retained unchanged in source_metadata records.
    """
    root = Path(root).expanduser().resolve(strict=True)
    if not root.is_dir():
        raise ValueError("BRACS root must be an existing directory")
    options = dict(metadata) if isinstance(metadata, dict) else {"path": str(metadata)}
    if not options.get("path"):
        raise ValueError("BRACS import requires the local XLSX/CSV metadata path")
    path = Path(options["path"]).expanduser().resolve(strict=True)
    rows, headers, sheet = _read_metadata(path, options.get("sheet"))
    columns = _resolve_columns(headers, options.get("column_mapping"))
    slides_dir = Path(options.get("slides_dir", root)).expanduser()
    slides_dir = (root / slides_dir if not slides_dir.is_absolute() else slides_dir).resolve(
        strict=True
    )
    if not slides_dir.is_dir() or not slides_dir.is_relative_to(root):
        raise ValueError("BRACS slides_dir must be a directory within root")
    pattern = options.get("glob", "**/*")
    if (
        not isinstance(pattern, str)
        or Path(pattern).is_absolute()
        or ".." in Path(pattern).parts
        or "\\" in pattern
    ):
        raise ValueError("BRACS glob must remain relative to slides_dir")
    files = {}
    for candidate in sorted(slides_dir.glob(pattern)):
        if not candidate.is_file() or candidate.suffix.lower() != ".svs":
            continue
        actual = candidate.resolve()
        if not actual.is_relative_to(root):
            raise ValueError("BRACS slide symlink targets must remain within root")
        files.setdefault(candidate.stem, []).append(candidate)

    records = []
    grouped = {}
    missing_identifier_records = []
    for index, row in enumerate(rows):
        subject = _text(row.get(columns["subject_id"])) if columns["subject_id"] else ""
        record_id = (
            f"BRACS:metadata:{index}:{_identity(json.dumps(row, sort_keys=True, allow_nan=False))}"
        )
        records.append(
            MetadataRecord(
                record_id, "source_metadata", [subject] if subject else [], row, str(path)
            )
        )
        value = _text(row.get(columns["slide_id"]))
        if not value:
            missing_identifier_records.append(record_id)
            continue
        if Path(value).name != value or "\\" in value or value in {".", ".."}:
            raise ValueError(
                "BRACS WSI identifiers must be explicit stems or SVS basenames, without directory components"
            )
        if Path(value).suffix and Path(value).suffix.lower() != ".svs":
            raise ValueError(
                "BRACS WSI identifier includes an unexpected extension; select the WSI metadata sheet"
            )
        stem = Path(value).stem if value.lower().endswith(".svs") else value
        grouped.setdefault(stem, []).append((row, record_id, subject))

    assets = []
    for stem in sorted(set(files) | set(grouped)):
        entries = grouped.get(stem, [])
        labels = sorted(
            {
                _text(row.get(columns["label"]))
                for row, _, _ in entries
                if columns["label"] and _text(row.get(columns["label"]))
            }
        )
        splits = sorted(
            {
                _text(row.get(columns["split"]))
                for row, _, _ in entries
                if columns["split"] and _text(row.get(columns["split"]))
            }
        )
        subjects = sorted({subject for _, _, subject in entries if subject})
        candidates = files.get(stem, [])
        conflict = len(labels) > 1 or len(splits) > 1 or len(subjects) > 1 or len(candidates) > 1
        status = "conflicting" if conflict else "known" if labels else "missing"
        if not candidates:
            candidates = [slides_dir / (stem + ".svs")]
        for candidate in candidates:
            relative = candidate.relative_to(root).as_posix()
            present = candidate.is_file()
            assets.append(
                Asset(
                    asset_id=f"BRACS:{_identity(relative)}",
                    dataset="BRACS",
                    kind="slide",
                    filename=candidate.name,
                    relative_path=relative,
                    size=candidate.stat().st_size if present else None,
                    slide_id=stem,
                    subject_ids=subjects,
                    labels={"wsi_label": labels[0]} if len(labels) == 1 else {},
                    label_status=status,
                    split=splits[0] if len(splits) == 1 else None,
                    metadata={
                        "source_record_ids": [record_id for _, record_id, _ in entries],
                        "observed_labels": labels,
                        "observed_splits": splits,
                        "local_status": "present" if present else "missing",
                        "duplicate_slide_stem": len(files.get(stem, [])) > 1,
                        "metadata_status": "present" if entries else "missing",
                    },
                )
            )
    return AcquisitionPlan(
        "BRACS",
        {"local_root": str(root), "metadata": _clean(options)},
        assets,
        records,
        {
            "source": SOURCE_URL,
            "metadata_sources": [str(path)],
            "imported_at": utc_now(),
            "acquisition": "authorized local import; SVS files registered in place",
            "metadata_sheet": sheet,
            "metadata_availability": {
                "labels": "source_workbook",
                "clinical": "not_provided_by_this_release",
            },
            "missing_identifier_records": missing_identifier_records,
            "roi_assets": "not_imported",
        },
    )


def discover(selection: dict, *, session=None) -> AcquisitionPlan:
    if isinstance(selection, dict) and selection.get("local_root") and selection.get("metadata"):
        return import_local(selection["local_root"], selection["metadata"])
    raise ValueError(
        "BRACS requires authorized local files: use import_local(root, metadata); register and download through the official BRACS portal first"
    )
