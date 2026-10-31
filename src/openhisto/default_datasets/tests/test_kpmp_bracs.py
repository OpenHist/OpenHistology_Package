"""Synthetic provider tests: no external requests or real dataset records."""

import csv
from datetime import date
import importlib.util
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openhisto.default_datasets.KPMP.provider import discover as discover_kpmp
from openhisto.default_datasets.KPMP.provider import import_local as import_kpmp
from openhisto.default_datasets.KPMP.provider import parse_clinical
from openhisto.default_datasets.BRACS.provider import discover as discover_bracs
from openhisto.default_datasets.BRACS.provider import import_local as import_bracs


ATLAS_HEADERS = [
    "Access",
    "Participant ID",
    "File UUID",
    "File Name",
    "Data Category",
    "Data Format",
    "Experimental Strategy",
    "Workflow Type",
    "Internal Package ID",
    "Size",
]


def write_csv(path, headers, rows):
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)
    return path


def atlas_row(**changes):
    row = dict(
        zip(
            ATLAS_HEADERS,
            [
                "Open",
                "SYNTH-SUBJECT-A",
                "synthetic-file-001",
                "synthetic-file-001_full slide.svs",
                "Histology",
                "SVS",
                "Light Microscopy",
                "Light Microscopy",
                "synthetic package/A",
                "2.5 GB",
            ],
        )
    )
    row.update(changes)
    return row


class KPMPProviderTests(unittest.TestCase):
    def setUp(self):
        self.temporary = TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)

    def repository(self, rows):
        return write_csv(self.root / "synthetic_repository.csv", ATLAS_HEADERS, rows)

    def test_export_controls_selection_preserves_full_filename_and_approximate_size(self):
        clinical = atlas_row(
            **{
                "File UUID": "synthetic-clinical",
                "File Name": "synthetic-clinical_binned.csv",
                "Data Category": "Clinical",
                "Data Format": "CSV",
                "Experimental Strategy": "Clinical data",
                "Workflow Type": "Binned public clinical",
                "Size": "9 KB",
            }
        )
        rows = [
            atlas_row(),
            atlas_row(**{"Access": "Controlled", "File UUID": "synthetic-controlled"}),
            clinical,
        ]
        path = self.repository(rows)
        plan = discover_kpmp({"repository_csv": path})
        self.assertEqual(len(plan.records), 3)
        self.assertEqual(len(plan.assets), 2)
        slide = next(asset for asset in plan.assets if asset.kind == "slide")
        self.assertEqual(slide.filename, rows[0]["File Name"])
        self.assertTrue(
            slide.url.endswith("synthetic%20package%2FA/synthetic-file-001_full%20slide.svs")
        )
        self.assertEqual(slide.relative_path, "slides/synthetic-file-001.svs")
        self.assertIsNone(slide.size)
        self.assertEqual(plan.records[0].data["Size"], "2.5 GB")
        self.assertEqual(plan.provenance["excluded_controlled_rows"], 1)
        self.assertEqual(plan.provenance["metadata_availability"]["clinical"], "public_binned")
        self.assertEqual(plan.provenance["metadata_sources"], [str(path.resolve())])
        json.dumps(plan.to_dict(), allow_nan=False)

    def test_missing_clinical_is_explicit_and_local_clinical_joins_by_participant(self):
        path = self.repository([atlas_row()])
        plan = discover_kpmp({"repository_csv": str(path)})
        self.assertEqual(
            plan.provenance["metadata_availability"]["clinical"], "not_in_repository_export"
        )
        clinical = write_csv(
            self.root / "synthetic_clinical.csv",
            ["Participant ID", "Age bin", "Unrecognized source field"],
            [
                {
                    "Participant ID": "SYNTH-SUBJECT-A",
                    "Age bin": "50–59",
                    "Unrecognized source field": "source value",
                },
                {
                    "Participant ID": "SYNTH-SUBJECT-B",
                    "Age bin": "60–69",
                    "Unrecognized source field": "another value",
                },
            ],
        )
        plan = discover_kpmp({"repository_csv": path, "clinical_path": clinical})
        records = [record for record in plan.records if record.kind == "clinical"]
        self.assertEqual(len(records), 2)
        self.assertEqual(plan.assets[0].clinical_refs, [records[0].record_id])
        self.assertEqual(records[0].data["Age bin"], "50–59")
        self.assertEqual(records[0].data["Unrecognized source field"], "source value")

    def test_default_visible_atlas_headers_and_official_microscopy_phrase(self):
        headers = [
            "Access",
            "Participant ID",
            "File Name",
            "Experimental Strategy",
            "Internal Package ID",
        ]
        row = {name: atlas_row()[name] for name in headers}
        row["Experimental Strategy"] = "Light Microscopic Whole Slide Images"
        path = write_csv(self.root / "default_visible_export.csv", headers, [row])
        plan = discover_kpmp({"repository_csv": path})
        self.assertEqual(len(plan.assets), 1)
        self.assertEqual(plan.assets[0].filename, row["File Name"])
        self.assertTrue(plan.assets[0].relative_path.endswith(".svs"))
        self.assertEqual(
            plan.assets[0].asset_id, discover_kpmp({"repository_csv": path}).assets[0].asset_id
        )
        self.assertIn("retrieved_at", plan.provenance)

    def test_json_jsonl_mapping_and_nonfinite_values(self):
        path = self.root / "synthetic_clinical.json"
        path.write_text(json.dumps([{"group": "SYNTH-A", "missing": float("nan")}]))
        records = parse_clinical(path, column_mapping={"subject_id": "group"})
        self.assertEqual(records[0].subject_ids, ["SYNTH-A"])
        self.assertIsNone(records[0].data["missing"])
        json.dumps(records[0].to_dict(), allow_nan=False)
        path = self.root / "synthetic_clinical.jsonl"
        path.write_text(
            '{"Participant ID":"SYNTH-A","bin":"x"}\n{"Participant ID":null,"bin":"y"}\n'
        )
        self.assertEqual(parse_clinical(path)[1].subject_ids, [])
        with self.assertRaisesRegex(ValueError, "Participant ID"):
            parse_clinical(self.root / "synthetic_clinical.json")

    def test_import_registers_full_filename_in_place_and_missing_files(self):
        row = atlas_row()
        repository = self.repository(
            [
                row,
                atlas_row(
                    **{"File UUID": "synthetic-missing", "File Name": "synthetic-missing.svs"}
                ),
            ]
        )
        slides = self.root / "authorized_local_slides"
        slides.mkdir()
        slide = slides / row["File Name"]
        slide.write_bytes(b"synthetic slide placeholder")
        plan = import_kpmp(self.root, repository)
        self.assertEqual(plan.assets[0].relative_path, slide.relative_to(self.root).as_posix())
        self.assertEqual(plan.assets[0].size, slide.stat().st_size)
        self.assertIsNone(plan.assets[0].url)
        self.assertEqual(plan.assets[1].metadata["local_status"], "missing")

    def test_rejects_incomplete_export_and_escaping_local_symlink(self):
        with self.assertRaisesRegex(ValueError, "repository_csv"):
            discover_kpmp({})
        incomplete = write_csv(
            self.root / "incomplete.csv", ["File Name"], [{"File Name": "synthetic.svs"}]
        )
        with self.assertRaisesRegex(ValueError, "Internal Package ID"):
            discover_kpmp({"repository_csv": incomplete})
        repository = self.repository([atlas_row()])
        with TemporaryDirectory() as external:
            outside = Path(external) / "synthetic_external.svs"
            outside.write_bytes(b"synthetic")
            (self.root / atlas_row()["File Name"]).symlink_to(outside)
            with self.assertRaisesRegex(ValueError, "within root"):
                import_kpmp(self.root, repository)


class BRACSProviderTests(unittest.TestCase):
    def setUp(self):
        self.temporary = TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)

    def test_preserves_all_rows_conflicts_missing_slides_and_unknown_fields(self):
        (self.root / "SYNTH_001.svs").write_bytes(b"synthetic slide")
        (self.root / "SYNTH_unlisted.svs").write_bytes(b"synthetic unlisted")
        (self.root / "SYNTH_001_roi.png").write_bytes(b"not imported")
        headers = ["WSI", "WSI label", "Patient ID", "Set", "Extra source field"]
        rows = [
            dict(zip(headers, ["SYNTH_001.svs", "PB", "SYNTH-GROUP-A", "Training", "first"])),
            dict(zip(headers, ["SYNTH_001", "IC", "SYNTH-GROUP-A", "Testing", "second"])),
            dict(zip(headers, ["SYNTH_002", "N", "SYNTH-GROUP-B", "Validation", "missing slide"])),
            dict(zip(headers, ["", "N", "SYNTH-GROUP-C", "Training", "missing identifier"])),
        ]
        path = write_csv(self.root / "synthetic_metadata.csv", headers, rows)
        plan = import_bracs(self.root, path)
        self.assertEqual(len(plan.records), 4)
        self.assertEqual(len(plan.assets), 3)
        assets = {asset.slide_id: asset for asset in plan.assets}
        conflict = assets["SYNTH_001"]
        self.assertEqual(conflict.label_status, "conflicting")
        self.assertEqual(conflict.labels, {})
        self.assertIsNone(conflict.split)
        self.assertEqual(conflict.metadata["observed_labels"], ["IC", "PB"])
        self.assertEqual(len(conflict.metadata["source_record_ids"]), 2)
        self.assertEqual(assets["SYNTH_002"].metadata["local_status"], "missing")
        self.assertEqual(assets["SYNTH_unlisted"].label_status, "missing")
        self.assertEqual(plan.records[1].data["Extra source field"], "second")
        self.assertEqual(len(plan.provenance["missing_identifier_records"]), 1)
        self.assertTrue(all(asset.kind == "slide" for asset in plan.assets))

    def test_unknown_identifier_requires_mapping_and_preserves_original_split(self):
        (self.root / "SYNTH_A.svs").write_bytes(b"synthetic")
        path = write_csv(
            self.root / "mapped.csv",
            ["custom_id", "custom_label", "custom_group", "custom_set"],
            [
                {
                    "custom_id": "SYNTH_A",
                    "custom_label": "ADH",
                    "custom_group": "SYNTH-P",
                    "custom_set": "Validation",
                },
            ],
        )
        with self.assertRaisesRegex(ValueError, "slide_id"):
            import_bracs(self.root, path)
        options = {
            "path": path,
            "column_mapping": {
                "slide_id": "custom_id",
                "label": "custom_label",
                "subject_id": "custom_group",
                "split": "custom_set",
            },
        }
        plan = import_bracs(self.root, options)
        self.assertEqual(plan.assets[0].split, "Validation")
        self.assertEqual(plan.assets[0].labels, {"wsi_label": "ADH"})
        self.assertEqual(plan.assets[0].subject_ids, ["SYNTH-P"])
        json.dumps(plan.to_dict(), allow_nan=False)

    def test_duplicate_files_are_flagged_and_root_escape_is_rejected(self):
        for directory in ("first", "second"):
            (self.root / directory).mkdir()
            (self.root / directory / "SYNTH_A.svs").write_bytes(b"synthetic")
        path = write_csv(
            self.root / "metadata.csv", ["WSI", "WSI label"], [{"WSI": "SYNTH_A", "WSI label": "N"}]
        )
        plan = import_bracs(self.root, path)
        self.assertEqual(len(plan.assets), 2)
        self.assertTrue(all(asset.label_status == "conflicting" for asset in plan.assets))
        with TemporaryDirectory() as external:
            with self.assertRaisesRegex(ValueError, "within root"):
                import_bracs(self.root, {"path": path, "slides_dir": external})

    def test_discover_is_actionable_without_unauthorized_downloads(self):
        with self.assertRaisesRegex(ValueError, "import_local"):
            discover_bracs({})

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "optional openpyxl unavailable")
    def test_generated_workbooks_and_clinical_dates_are_json_safe(self):
        import openpyxl

        (self.root / "SYNTH_A.svs").write_bytes(b"synthetic")
        path = self.root / "synthetic_bracs.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active.title = "Unrelated"
        sheet = workbook.create_sheet("WSI_Information")
        sheet.append(["WSI", "WSI label", "Patient ID", "Set"])
        sheet.append(["SYNTH_A", "FEA", "SYNTH-P", "Training"])
        workbook.save(path)
        workbook.close()
        plan = import_bracs(self.root, path)
        self.assertEqual(plan.assets[0].labels, {"wsi_label": "FEA"})
        self.assertEqual(plan.provenance["metadata_sheet"], "WSI_Information")

        path = self.root / "synthetic_clinical.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active.append(["Participant ID", "Synthetic date"])
        workbook.active.append(["SYNTH-P", date(2000, 1, 2)])
        workbook.save(path)
        workbook.close()
        records = parse_clinical(path)
        self.assertTrue(records[0].data["Synthetic date"].startswith("2000-01-02"))
        json.dumps(records[0].to_dict(), allow_nan=False)


if __name__ == "__main__":
    unittest.main()
